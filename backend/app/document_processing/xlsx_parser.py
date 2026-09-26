"""
Excel (.xlsx) -> ParsedDocument (Phase 16).

Each visible worksheet becomes one "page" (page 2 = the second sheet), opened
by a heading "2. <sheet name>" so every chunk and citation names its sheet.
Rows become tables. A long sheet is split into several tables, each small
enough for one chunk and each repeating the header row, so a chunk of rows
500-540 still says which column is which.

Values are the ones Excel last calculated (formulas are not re-evaluated),
formatted as the user would read them: dates as ISO dates, whole numbers
without ".0". Empty rows and trailing empty columns are dropped.

Limits: `max_cells` bounds the work (and the embedding cost) of a huge
workbook; what is left out is reported as a warning. Chart sheets, hidden
sheets, comments and formatting are not content and are not read.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook

from app.document_processing.parser import (
    BlockKind,
    ParsedDocument,
    ParsedPage,
    Table,
    TextBlock,
)

PAGE_WIDTH = 612.0
LINE = 14.0
# A table's Markdown stays under about one child chunk (chunk_max_tokens=400
# at ~4 characters per token), header included.
MAX_TABLE_CHARS = 1400
MAX_CELL_CHARS = 200


@dataclass
class XlsxParseResult:
    document: ParsedDocument
    warnings: list[str] = field(default_factory=list)


def format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    text = " ".join(str(value).split())
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def _row_chars(row: list[str]) -> int:
    return sum(len(c) + 3 for c in row) + 2


def split_rows(header: list[str], body: list[list[str]]) -> list[list[list[str]]]:
    """Group body rows into tables of at most MAX_TABLE_CHARS, each starting
    with the header. A single very wide row still gets a table of its own."""
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    size = _row_chars(header)
    for row in body:
        if current and size + _row_chars(row) > MAX_TABLE_CHARS:
            tables.append([header, *current])
            current, size = [], _row_chars(header)
        current.append(row)
        size += _row_chars(row)
    if current or not tables:
        tables.append([header, *current])
    return tables


def parse_xlsx(data: bytes, *, max_cells: int) -> XlsxParseResult:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    warnings: list[str] = []
    pages: list[ParsedPage] = []
    cells_left = max_cells
    try:
        sheets = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        for ws in sheets:
            if cells_left <= 0:
                warnings.append(
                    f"Cell limit reached: sheet '{ws.title}' and later sheets were not read."
                )
                break
            rows: list[list[str]] = []
            truncated = False
            for raw in ws.iter_rows(values_only=True):
                row = [format_cell(v) for v in raw]
                while row and not row[-1]:
                    row.pop()
                if not row:
                    continue
                if len(row) > cells_left:
                    truncated = True
                    break
                cells_left -= len(row)
                rows.append(row)
            if truncated:
                warnings.append(f"Cell limit reached: sheet '{ws.title}' was read only in part.")
            pages.append(_sheet_page(len(pages) + 1, ws.title, rows))
    finally:
        workbook.close()

    if not pages:
        warnings.append("The workbook has no visible sheets with data.")
        pages.append(ParsedPage(number=1, width=PAGE_WIDTH, height=LINE))
    metadata = {"format": "xlsx", "sheets": [p.blocks[0].text for p in pages if p.blocks]}
    return XlsxParseResult(ParsedDocument(pages=pages, metadata=metadata), warnings)


def _sheet_page(number: int, title: str, rows: list[list[str]]) -> ParsedPage:
    page = ParsedPage(number=number, width=PAGE_WIDTH, height=LINE)
    # Numbered, so the chunker treats it as a section ("2. Invoices") rather
    # than taking the first sheet's name as the document title.
    page.blocks.append(
        TextBlock(
            text=f"{number}. {title}",
            bbox=(0.0, 0.0, PAGE_WIDTH, LINE),
            kind=BlockKind.HEADING,
            font_size=16.0,
            is_bold=True,
        )
    )
    if not rows:
        page.blocks.append(TextBlock(text="(empty sheet)", bbox=(0.0, LINE, PAGE_WIDTH, 2 * LINE)))
        return page
    width = max(len(r) for r in rows)
    grid = [r + [""] * (width - len(r)) for r in rows]
    header, body = grid[0], grid[1:]
    y = LINE
    first_row = 1
    for index, table_rows in enumerate(split_rows(header, body)):
        y += LINE
        table = Table(bbox=(0.0, y, PAGE_WIDTH, y + LINE), rows=table_rows)
        last_row = first_row + len(table_rows) - 2
        if index > 0:
            # Only a sheet's first table is summarised by the model (Phase 9);
            # the rest say where they sit, which costs nothing.
            table.summary = f"Sheet '{title}', data rows {first_row}-{last_row} (continued)."
        page.tables.append(table)
        first_row = last_row + 1
        y += LINE * len(table_rows)
    page.height = y
    return page
