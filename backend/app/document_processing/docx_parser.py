"""
Word (.docx) -> ParsedDocument (Phase 16).

A .docx has no fixed pages or coordinates, so it is read as a structure
instead: headings (from paragraph styles), paragraphs, list items and
tables, in document order. The result has the same shape as a parsed PDF,
so chunking, embeddings, search and citations work unchanged.

Pages
-----
Page numbers in citations come from the page breaks Word records in the
file: explicit breaks (Ctrl+Enter, "page break before") and the
`lastRenderedPageBreak` markers Word writes where pages ended when the file
was last saved. Files produced by other tools may have neither; their
content is all "page 1", and citations still carry the section heading.

Coordinates
-----------
Blocks get synthetic, increasing y positions (one "line" each), which is
all the chunker needs: it orders tables and text by `bbox[1]`.

Headers, footers, comments and tracked deletions are not content and are
not read. Embedded pictures are returned for captioning like PDF images.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from docx import Document as open_docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from PIL import Image

from app.document_processing.parser import (
    BlockKind,
    ParsedDocument,
    ParsedPage,
    Table,
    TextBlock,
)

PAGE_WIDTH = 612.0  # US Letter in points; nominal, for a consistent shape
LINE = 14.0  # synthetic height of one block


@dataclass
class DocxImage:
    """An embedded picture, with the page it sits on."""

    page_number: int
    y: float  # position among the page's blocks, so it is chunked in place
    png: bytes
    width_px: int
    height_px: int


@dataclass
class DocxParseResult:
    document: ParsedDocument
    images: list[DocxImage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _heading_level(paragraph: Paragraph) -> int | None:
    """1 for Title/Heading 1, 2 for Heading 2, ...; None for body text."""
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if name == "Title":
        return 1
    if name.startswith("Heading"):
        digits = "".join(ch for ch in name if ch.isdigit())
        return int(digits) if digits else 1
    return None


def _is_list_item(paragraph: Paragraph) -> bool:
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    ppr = paragraph._p.pPr
    return name.startswith("List") or (ppr is not None and ppr.numPr is not None)


def _breaks_before(paragraph: Paragraph) -> int:
    """Page breaks that occur at the start of, or inside, this paragraph."""
    p = paragraph._p
    count = len(p.findall(".//" + qn("w:lastRenderedPageBreak")))
    count += sum(1 for br in p.iter(qn("w:br")) if br.get(qn("w:type")) == "page")
    ppr = p.pPr
    if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None:
        count += 1
    return count


def _table_rows(table: DocxTable) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            # Merged cells are returned once per grid column; keep one copy.
            if cell._tc is previous:
                continue
            previous = cell._tc
            cells.append(" ".join(cell.text.split()))
        if any(cells):
            rows.append(cells)
    return rows


def _picture_ids(paragraph: Paragraph) -> list[str]:
    return [
        blip.get(qn("r:embed"))
        for blip in paragraph._p.iter(qn("a:blip"))
        if blip.get(qn("r:embed"))
    ]


def parse_docx(
    data: bytes, *, max_pages: int, max_images: int, min_image_dimension_px: int
) -> DocxParseResult:
    document: DocxDocument = open_docx(io.BytesIO(data))
    warnings: list[str] = []
    pages: list[ParsedPage] = [ParsedPage(number=1, width=PAGE_WIDTH, height=0.0)]
    images: list[DocxImage] = []
    seen_images: set[str] = set()
    y = 0.0
    page_limit_hit = False

    def page() -> ParsedPage:
        return pages[-1]

    def next_y() -> float:
        nonlocal y
        y += LINE
        return y

    body = document.element.body
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            # Only break before content: a break at the very start of the
            # document (or a run of breaks) must not create empty pages.
            for _ in range(_breaks_before(paragraph)):
                if page().blocks or page().tables or page().images:
                    if len(pages) >= max_pages:
                        page_limit_hit = True
                        break
                    pages.append(ParsedPage(number=len(pages) + 1, width=PAGE_WIDTH, height=0.0))
                    y = 0.0
            for rel_id in _picture_ids(paragraph):
                if rel_id in seen_images or len(images) >= max_images:
                    continue
                seen_images.add(rel_id)
                picture = _picture_png(document, rel_id, min_image_dimension_px)
                if picture is not None:
                    png, width, height = picture
                    images.append(DocxImage(page().number, next_y(), png, width, height))
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            level = _heading_level(paragraph)
            top = next_y()
            if level is not None:
                block = TextBlock(
                    text=text,
                    bbox=(0.0, top, PAGE_WIDTH, top + LINE),
                    kind=BlockKind.HEADING,
                    font_size=max(10.0, 20.0 - 2 * level),
                    is_bold=True,
                )
            else:
                kind = BlockKind.LIST_ITEM if _is_list_item(paragraph) else BlockKind.PARAGRAPH
                block = TextBlock(text=text, bbox=(0.0, top, PAGE_WIDTH, top + LINE), kind=kind)
            page().blocks.append(block)
        elif element.tag == qn("w:tbl"):
            rows = _table_rows(DocxTable(element, document))
            if rows:
                top = next_y()
                page().tables.append(Table(bbox=(0.0, top, PAGE_WIDTH, top + LINE), rows=rows))
                y += LINE * len(rows)

    if page_limit_hit:
        warnings.append(
            f"The document has more than {max_pages} pages; later content is cited as "
            f"page {max_pages}."
        )
    for p in pages:
        p.height = max(y, LINE)
    properties = document.core_properties
    metadata = {
        "format": "docx",
        "title": properties.title or None,
        "author": properties.author or None,
    }
    return DocxParseResult(ParsedDocument(pages=pages, metadata=metadata), images, warnings)


def _picture_png(
    document: DocxDocument, rel_id: str, min_dimension_px: int
) -> tuple[bytes, int, int] | None:
    """The picture as PNG, or None if it is small (an icon), not a raster
    image (e.g. EMF/WMF vector art) or cannot be decoded."""
    part = document.part.related_parts.get(rel_id)
    blob = getattr(part, "blob", None)
    if not blob:
        return None
    try:
        with Image.open(io.BytesIO(blob)) as image:
            width, height = image.size
            if min(width, height) < min_dimension_px:
                return None
            out = io.BytesIO()
            image.convert("RGB").save(out, "PNG")
            return out.getvalue(), width, height
    except Exception:  # noqa: BLE001 — one bad picture must not fail the document
        return None
