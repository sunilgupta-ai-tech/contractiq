"""Phase 16: Word, Excel and image uploads — format detection, parsing into
the shared ParsedDocument, and handwriting transcription helpers.

Every file is generated in the test, so nothing binary is checked in."""

import io
import zipfile
from datetime import datetime

import openpyxl
import pymupdf
import pytest
from docx import Document as new_docx
from docx.enum.text import WD_BREAK
from PIL import Image

from app.chunking.structure import build_segments
from app.core.exceptions import InvalidFileError
from app.db.models import FileType
from app.document_processing.docx_parser import parse_docx
from app.document_processing.formats import detect_format, format_for_storage_key
from app.document_processing.image_input import image_to_pdf, image_to_png
from app.document_processing.parser import (
    BlockKind,
    ParsedDocument,
    ParsedPage,
    TextBlock,
    TextSource,
)
from app.document_processing.xlsx_parser import (
    MAX_TABLE_CHARS,
    format_cell,
    parse_xlsx,
    split_rows,
)
from app.multimodal.transcriber import is_weak_ocr, transcript_to_page_content
from app.services.document_service import validate_upload

# --- Generated files --------------------------------------------------------------


def make_docx() -> bytes:
    doc = new_docx()
    doc.add_heading("Leave Policy", level=0)
    doc.add_heading("1. Annual leave", level=1)
    doc.add_paragraph("Employees receive 24 days of paid leave per year.")
    doc.add_paragraph("Carry-over is capped at 5 days.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Grade", "Days"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Manager", "28"
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("2. Sick leave", level=1)
    doc.add_paragraph("Up to 10 days with a medical certificate.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_xlsx(rows: int = 3) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Invoice", "Vendor", "Amount", "Due"])
    for i in range(rows):
        ws.append([f"INV-{i + 1}", "Acme", 1200.0 + i, datetime(2026, 10, i % 28 + 1)])
    hidden = wb.create_sheet("Secret")
    hidden.append(["not", "indexed"])
    hidden.sheet_state = "hidden"
    wb.create_sheet("Empty")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_image(fmt: str = "PNG", size: tuple[int, int] = (400, 300)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, "white").save(out, fmt)
    return out.getvalue()


def zip_with(names: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, content in names.items():
            archive.writestr(name, content)
    return out.getvalue()


# --- Detection --------------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "data", "file_type", "stored"),
    [
        ("MSA.PDF", b"%PDF-1.7\n", FileType.PDF, "original.pdf"),
        ("scan.jpeg", make_image("JPEG"), FileType.IMAGE, "original.jpg"),
        ("receipt.png", make_image("PNG"), FileType.IMAGE, "original.png"),
        ("policy.docx", make_docx(), FileType.WORD, "original.docx"),
        ("data.xlsx", make_xlsx(), FileType.EXCEL, "original.xlsx"),
    ],
)
def test_supported_formats_are_detected(filename, data, file_type, stored):
    fmt = validate_upload(filename, "application/octet-stream", data)
    assert (fmt.file_type, fmt.stored_filename) == (file_type, stored)
    assert format_for_storage_key(f"tenants/t/documents/d/v/{stored}") == fmt


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "save it as .docx"),
        ("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "save it as .xlsx"),
        ("macro.xlsm", make_xlsx(), "Macro-enabled"),
        ("locked.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "password-protected"),
        ("tool.exe", b"MZ\x90\x00", "not supported"),
        ("renamed.docx", b"%PDF-1.7\n", "not a valid Word"),
        ("renamed.xlsx", make_docx(), "not a valid Excel"),
        ("renamed.png", make_image("JPEG"), "not a valid PNG"),
        ("broken.jpg", b"\xff\xd8\xff\xe0garbage", "damaged"),
        ("broken.docx", b"PK\x03\x04garbage", "damaged"),
        (
            "macros.docx",
            zip_with(
                {"[Content_Types].xml": b"", "word/document.xml": b"", "word/vbaProject.bin": b""}
            ),
            "macros",
        ),
        ("empty.png", b"", "empty"),
    ],
)
def test_unsupported_or_disguised_files_are_refused(filename, data, message):
    with pytest.raises(InvalidFileError, match=message):
        detect_format(filename, data)


def test_html_declared_type_is_refused():
    with pytest.raises(InvalidFileError, match="not supported"):
        validate_upload("page.pdf", "text/html", b"%PDF-1.7\n")


def test_legacy_versions_stored_as_pdf():
    assert format_for_storage_key("tenants/t/documents/d/v/original.pdf").file_type is FileType.PDF
    assert format_for_storage_key("unknown").file_type is FileType.PDF


# --- Word ---------------------------------------------------------------------------


def test_docx_becomes_headings_lists_tables_and_pages():
    result = parse_docx(make_docx(), max_pages=100, max_images=10, min_image_dimension_px=150)
    doc = result.document
    assert doc.page_count == 2
    first = doc.pages[0]
    kinds = [(b.kind, b.text) for b in first.blocks]
    assert kinds[:3] == [
        (BlockKind.HEADING, "Leave Policy"),
        (BlockKind.HEADING, "1. Annual leave"),
        (BlockKind.PARAGRAPH, "Employees receive 24 days of paid leave per year."),
    ]
    assert (BlockKind.LIST_ITEM, "Carry-over is capped at 5 days.") in kinds
    assert first.tables[0].rows == [["Grade", "Days"], ["Manager", "28"]]
    assert doc.pages[1].blocks[0].text == "2. Sick leave"
    # Chunking sees sections and the table in order.
    title, segments = build_segments(doc)
    assert title == "Leave Policy"
    assert [s.section for s in segments if not s.is_table][:1] == ["1"]
    assert any(s.is_table for s in segments)
    assert segments[-1].section == "2"


def test_docx_round_trips_through_parsed_json():
    doc = parse_docx(make_docx(), max_pages=5, max_images=0, min_image_dimension_px=150).document
    again = ParsedDocument.from_dict(doc.to_dict())
    assert [p.text for p in again.pages] == [p.text for p in doc.pages]


# --- Excel --------------------------------------------------------------------------


def test_xlsx_sheets_become_pages_with_header_tables():
    doc = parse_xlsx(make_xlsx(), max_cells=10_000).document
    # Hidden sheets are skipped; the empty sheet is still listed.
    assert [p.blocks[0].text for p in doc.pages] == ["1. Invoices", "2. Empty"]
    table = doc.pages[0].tables[0]
    assert table.rows[0] == ["Invoice", "Vendor", "Amount", "Due"]
    assert table.rows[1] == ["INV-1", "Acme", "1200", "2026-10-01"]
    assert table.summary is None  # the first table of a sheet is summarised by the model
    _, segments = build_segments(doc)
    assert segments[0].section == "1" and segments[0].section_title == "Invoices"


def test_long_sheets_split_into_small_tables_repeating_the_header():
    doc = parse_xlsx(make_xlsx(rows=400), max_cells=100_000).document
    tables = doc.pages[0].tables
    assert len(tables) > 1
    assert all(t.rows[0] == ["Invoice", "Vendor", "Amount", "Due"] for t in tables)
    assert sum(len(t.rows) - 1 for t in tables) == 400
    assert all(len(t.to_markdown()) <= MAX_TABLE_CHARS + 200 for t in tables)
    # Continuation tables describe themselves, so no model call is needed.
    assert tables[1].summary.startswith("Sheet 'Invoices', data rows ")


def test_cell_limit_is_reported():
    result = parse_xlsx(make_xlsx(rows=50), max_cells=40)
    assert any("Cell limit" in w for w in result.warnings)


@pytest.mark.parametrize(
    ("value", "text"),
    [
        (None, ""),
        (3.0, "3"),
        (2.5, "2.5"),
        (True, "TRUE"),
        (datetime(2026, 1, 2), "2026-01-02"),
        (datetime(2026, 1, 2, 9, 30), "2026-01-02 09:30:00"),
        ("  a\nb  ", "a b"),
    ],
)
def test_cells_read_as_people_see_them(value, text):
    assert format_cell(value) == text


def test_split_rows_keeps_at_least_one_row_per_table():
    header = ["h"]
    wide = ["x" * (MAX_TABLE_CHARS * 2)]
    assert split_rows(header, [wide, wide]) == [[header, wide], [header, wide]]
    assert split_rows(header, []) == [[header]]


# --- Images -------------------------------------------------------------------------


def test_photos_are_upright_bounded_and_wrapped_as_a_scanned_page():
    photo = io.BytesIO()
    image = Image.new("RGB", (800, 400), "white")
    exif = image.getexif()
    exif[0x0112] = 6  # "rotate 90° clockwise to view"
    image.save(photo, "JPEG", exif=exif)
    png, width, height = image_to_png(photo.getvalue(), max_side_px=300)
    assert (width, height) == (150, 300)  # rotated upright, long side bounded

    pdf = image_to_pdf(png, width, height, dpi=300)
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        assert doc.page_count == 1
        page = doc[0]
        assert not page.get_text().strip()  # no text layer: OCR will run
        assert round(page.rect.width) == round(width * 72 / 300)


# --- Transcription ------------------------------------------------------------------


def _page(**kw) -> ParsedPage:
    return ParsedPage(number=1, width=600.0, height=800.0, **kw)


def test_weak_ocr_detection():
    words = " ".join(["word"] * 20)
    good = _page(blocks=[TextBlock(words, (0, 0, 1, 1))], is_scanned=True, ocr_confidence=91.0)
    faint = _page(blocks=[TextBlock(words, (0, 0, 1, 1))], is_scanned=True, ocr_confidence=41.0)
    empty = _page(is_scanned=True)
    digital = _page(blocks=[], is_scanned=False)
    assert not is_weak_ocr(good, min_confidence=70)
    assert is_weak_ocr(faint, min_confidence=70)
    assert is_weak_ocr(empty, min_confidence=70)
    assert not is_weak_ocr(digital, min_confidence=70)


def test_transcript_becomes_text_and_register_tables():
    transcript = """Stock register — March

| Date | Item | Qty |
| --- | --- | --- |
| 03/03 | Cement | 40 |
| 04/03 | Steel [illegible] | 12 |

- Checked by R. Mehta
Signed on 5 March."""
    blocks, tables = transcript_to_page_content(transcript, _page())
    assert [b.text for b in blocks] == [
        "Stock register — March",
        "- Checked by R. Mehta Signed on 5 March.",
    ]
    assert blocks[1].kind is BlockKind.LIST_ITEM
    assert all(b.source is TextSource.VISION for b in blocks)
    assert tables[0].rows == [
        ["Date", "Item", "Qty"],
        ["03/03", "Cement", "40"],
        ["04/03", "Steel [illegible]", "12"],
    ]
    # Reading order is preserved for the chunker.
    assert blocks[0].bbox[1] < tables[0].bbox[1] < blocks[1].bbox[1]
